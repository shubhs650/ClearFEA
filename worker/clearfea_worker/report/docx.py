# SPDX-License-Identifier: AGPL-3.0-or-later
"""Lays a ReportDocument out as a Word file (M7.3, ADR-077).

A4 portrait in one column, so the PDF of M7.4 can follow it page for page: a cover with the logo,
company name, title and title block; then numbered sections with numbered tables and figures; the
title, document number and revision in every page header after the cover and "Page n of N" in
every footer. Only python-docx's own default template is used, with Arial text.
"""

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

from .content import Figure, Heading, Paragraph, ReportDocument, Table

#: Page and print sizes.
PAGE_WIDTH, PAGE_HEIGHT = Mm(210), Mm(297)
MARGIN = Mm(20)
TEXT_WIDTH_MM = 170
FIGURE_WIDTH_MM = 160
#: Tallest a figure may print, so a figure and its caption fit on a page with its heading.
FIGURE_MAX_HEIGHT_MM = 200
LOGO_HEIGHT_MM = 18
LOGO_MAX_WIDTH_MM = 70

#: Tables of up to this many rows are kept on one page.
KEEP_TOGETHER_ROWS = 25

FONT = "Arial"
BODY_PT = 10
TABLE_PT = 9
HEADER_FILL = "E8ECF2"
NOTE_FILL = "F4F5F7"
DARK = RGBColor(0x1F, 0x2A, 0x3A)


def write_docx(document: ReportDocument, files: Path, path: Path, generated_at: str) -> int:
    """Writes the report to `path`, reading figures and the logo from `files`. Returns how many
    figures it placed."""
    doc = Document()
    _styles(doc)
    _page(doc)
    _properties(doc, document, generated_at)
    _cover(doc, document, files)
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.different_first_page_header_footer = False
    _header_footer(body, document)
    tables = figures = 0
    for number, section in enumerate(document.sections, start=1):
        doc.add_heading(f"{number}  {section.title}", level=1)
        for block in section.blocks:
            if isinstance(block, Heading):
                doc.add_heading(block.text, level=2)
            elif isinstance(block, Paragraph):
                _paragraph(doc, block)
            elif isinstance(block, Table):
                tables += 1
                _table(doc, block, tables)
            elif isinstance(block, Figure):
                figures += 1
                _figure(doc, block, figures, files)
    doc.save(str(path))
    return figures


def _styles(doc: Document) -> None:  # type: ignore[valid-type]
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    _east_asian(normal.element, FONT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, before in (("Heading 1", 14, 14), ("Heading 2", 11, 10), ("Title", 22, 0)):
        style = styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = DARK
        _east_asian(style.element, FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
    caption = styles["Caption"]
    caption.font.name = FONT
    caption.font.size = Pt(9)
    caption.font.bold = False
    caption.font.italic = False
    caption.font.color.rgb = DARK
    caption.paragraph_format.space_after = Pt(8)


def _east_asian(element, font: str) -> None:
    """Sets the font for every script, so Word does not swap in another for symbols like σ."""
    rpr = element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), font)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)


def _page(doc: Document) -> None:  # type: ignore[valid-type]
    section = doc.sections[0]
    section.page_width, section.page_height = PAGE_WIDTH, PAGE_HEIGHT
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, MARGIN)
    section.header_distance = section.footer_distance = Mm(10)


def _properties(doc: Document, document: ReportDocument, generated_at: str) -> None:  # type: ignore[valid-type]
    props = doc.core_properties
    moment = datetime.fromisoformat(generated_at.replace("Z", "+00:00")).replace(tzinfo=None)
    props.title = document.title
    props.subject = "Finite element analysis report"
    props.author = document.company or "ClearFEA"
    props.last_modified_by = "ClearFEA"
    props.comments = "Made by ClearFEA from a solved run."
    props.keywords = ""
    props.category = ""
    props.revision = 1
    props.created = props.modified = moment
    props.last_printed = moment


def _cover(doc: Document, document: ReportDocument, files: Path) -> None:  # type: ignore[valid-type]
    if document.logo is not None and document.logo_size is not None:
        width, height = document.logo_size
        logo_height = min(LOGO_HEIGHT_MM, LOGO_MAX_WIDTH_MM * height / width)
        doc.add_paragraph().add_run().add_picture(
            str(files / document.logo), height=Mm(logo_height)
        )
    if document.company:
        company = doc.add_paragraph()
        run = company.add_run(document.company)
        run.font.size = Pt(12)
        run.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(90)
    doc.add_paragraph(document.title, style="Title")
    subtitle = doc.add_paragraph()
    subtitle.add_run("Finite element analysis report").font.size = Pt(13)
    subtitle.paragraph_format.space_after = Pt(30)
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = doc.styles["Table Grid"]
    for line in document.cover:
        cells = table.add_row().cells
        _cell(cells[0], line.label, bold=True, fill=HEADER_FILL)
        _cell(cells[1], line.value)
    _widths(table, [1, 2], 120)


def _header_footer(section, document: ReportDocument) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.text = ""
    run = header.add_run(document.header)
    run.font.size = Pt(8)
    run.font.color.rgb = DARK
    _rule(header, "bottom")
    footer = section.footer.paragraphs[0]
    footer.text = ""
    # The Footer style's own centre stop would catch the tab before the right one.
    for stop in (Mm(82.55), Mm(165.1)):
        footer.paragraph_format.tab_stops.add_tab_stop(stop, WD_TAB_ALIGNMENT.CLEAR)
    footer.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_WIDTH_MM), WD_TAB_ALIGNMENT.RIGHT)
    left = document.company + "  ·  " if document.company else ""
    footer.add_run(f"{left}Made with ClearFEA\tPage ").font.size = Pt(8)
    _field(footer, "PAGE")
    footer.add_run(" of ").font.size = Pt(8)
    _field(footer, "NUMPAGES")
    _rule(footer, "top")


def _field(paragraph: DocxParagraph, code: str) -> None:
    """A Word field such as PAGE, with a placeholder value until Word or LibreOffice updates it."""

    def part(kind: str, text: str | None = None) -> None:
        run = paragraph.add_run()
        run.font.size = Pt(8)
        if kind in ("begin", "separate", "end"):
            char = OxmlElement("w:fldChar")
            char.set(qn("w:fldCharType"), kind)
            run._r.append(char)
        else:
            element = OxmlElement(kind)
            element.set(qn("xml:space"), "preserve")
            element.text = text
            run._r.append(element)

    part("begin")
    part("w:instrText", f" {code} ")
    part("separate")
    part("w:t", "1")
    part("end")


def _rule(paragraph: DocxParagraph, side: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    line = OxmlElement(f"w:{side}")
    for key, value in (("w:val", "single"), ("w:sz", "4"), ("w:space", "4"), ("w:color", "8A94A6")):
        line.set(qn(key), value)
    borders.append(line)
    ppr.append(borders)


def _paragraph(doc: Document, block: Paragraph) -> None:  # type: ignore[valid-type]
    paragraph = doc.add_paragraph(block.text)
    if block.style == "note":
        _shade(paragraph, NOTE_FILL)
        paragraph.paragraph_format.left_indent = Mm(2)
        paragraph.paragraph_format.right_indent = Mm(2)
        paragraph.paragraph_format.space_after = Pt(6)


def _shade(paragraph: DocxParagraph, fill: str) -> None:
    shading = OxmlElement("w:shd")
    for key, value in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", fill)):
        shading.set(qn(key), value)
    paragraph._p.get_or_add_pPr().append(shading)


def _table(doc: Document, block: Table, number: int) -> None:  # type: ignore[valid-type]
    caption = doc.add_paragraph(f"Table {number}. {block.caption}", style="Caption")
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(2)
    table = doc.add_table(rows=1, cols=len(block.columns))
    table.style = doc.styles["Table Grid"]
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, block.columns, strict=True):
        _cell(cell, text, bold=True, fill=HEADER_FILL)
    _repeat_header(table)
    for row in block.rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row, strict=True):
            _cell(cell, text)
    if len(block.rows) <= KEEP_TOGETHER_ROWS:
        # Each row kept with the next keeps a short table on one page.
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    _widths(table, block.widths or [1] * len(block.columns), TEXT_WIDTH_MM)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _cell(cell, text: str, bold: bool = False, fill: str | None = None) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.font.size = Pt(TABLE_PT)
    run.bold = bold
    if fill is not None:
        shading = OxmlElement("w:shd")
        for key, value in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", fill)):
            shading.set(qn(key), value)
        cell._tc.get_or_add_tcPr().append(shading)


def _repeat_header(table: DocxTable) -> None:
    trpr = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trpr.append(header)


def _widths(table: DocxTable, weights: list[float], total_mm: float) -> None:
    """Fixed column widths in proportion to `weights`, set on every cell as Word expects."""
    table.autofit = False
    total = sum(weights)
    widths = [Mm(total_mm * w / total) for w in weights]
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = width
    for column, width in zip(table.columns, widths, strict=True):
        column.width = width


def _figure(doc: Document, block: Figure, number: int, files: Path) -> None:  # type: ignore[valid-type]
    width = float(FIGURE_WIDTH_MM)
    if width * block.height / block.width > FIGURE_MAX_HEIGHT_MM:
        width = FIGURE_MAX_HEIGHT_MM * block.width / block.height
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_with_next = True
    picture.paragraph_format.space_before = Pt(6)
    picture.add_run().add_picture(str(files / block.file), width=Mm(width))
    caption = doc.add_paragraph(f"Figure {number}. {block.caption}", style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
