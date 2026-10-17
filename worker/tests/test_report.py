# SPDX-License-Identifier: AGPL-3.0-or-later
"""The report's content and its Word layout (M7.3, ADR-077)."""

import copy
import re
import struct
import zlib
from pathlib import Path
from typing import Any

import pytest
from docx import Document

from clearfea_worker.report.content import (
    Figure,
    Paragraph,
    Table,
    build_report,
    report_date,
    solver_warnings,
)
from clearfea_worker.report.docx import write_docx
from clearfea_worker.report.labels import (
    ANALYSIS_KINDS,
    LOAD_KINDS,
    SECTION_TITLES,
    SOLVER_LABELS,
    SUPPORT_KINDS,
)
from clearfea_worker.spec.geometry_schema import Topology
from clearfea_worker.spec.report_schema import ReportJobParams

REPO = Path(__file__).resolve().parents[2]

SECTIONS = [
    "project",
    "geometry",
    "materials",
    "mesh",
    "loadsSupports",
    "solver",
    "results",
    "reactions",
    "extremes",
    "warnings",
]


def png(width: int = 16, height: int = 10) -> bytes:
    """A small white PNG."""
    raw = b"".join(b"\x00" + b"\xff\xff\xff" * width for _ in range(height))

    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def topology() -> Topology:
    faces = [
        {
            "tag": tag,
            "type": "plane",
            "gmshType": "Plane",
            "area": 2e-4,
            "centroid": [0.0, 0.0, 0.0],
            "bbox": [0, 0, 0, 0.04, 0.02, 0.01],
            "bodies": [1],
        }
        for tag in range(1, 7)
    ]
    return Topology.model_validate(
        {
            "version": 1,
            "size": 0.0458,
            "bbox": [0, 0, 0, 0.04, 0.02, 0.01],
            "unit": {
                "declared": "mm",
                "detectedUnit": "mm",
                "problem": None,
                "readingFactor": 0.001,
            },
            "scaling": 1,
            "software": {"gmsh": "4.15.2", "occt": "7.9.2"},
            "bodies": [{"tag": 1, "name": "Block", "volume": 8e-6, "faces": [1, 2, 3, 4, 5, 6]}],
            "faces": faces,
        }
    )


def static_params(**changes: Any) -> dict[str, Any]:
    params: dict[str, Any] = {
        "runId": "run-1",
        "generatedAt": "2026-10-17T04:10:00.000Z",
        "unitSystem": "metric-mm",
        "sections": SECTIONS,
        "titleBlock": {
            "title": "Block under pressure",
            "documentNumber": "CF-001",
            "revision": "A",
            "preparedBy": "A. Engineer",
            "checkedBy": "",
            "approvedBy": "",
        },
        "companyName": "Example Engineering",
        "logo": {
            "contentType": "image/png",
            "sizeBytes": 100,
            "width": 16,
            "height": 10,
            "updatedAt": "2026-10-17T03:00:00.000Z",
        },
        "projectName": "My project",
        "simulationName": "Block",
        "geometry": {"name": "block.step", "unit": "mm"},
        "run": {
            "id": "run-1",
            "simulationId": "sim-1",
            "meshId": "mesh-1",
            "spec": {
                "version": 1,
                "analysis": "linear_static",
                "geometryId": "geo-1",
                "bodies": [{"tag": 1, "material": "m1"}],
                "materials": [
                    {
                        "id": "m1",
                        "name": "Steel",
                        "model": "isotropic_elastic",
                        "youngsModulus": {"value": 200, "unit": "GPa"},
                        "poissonsRatio": 0.3,
                        "yieldStrength": {"value": 36, "unit": "ksi"},
                    }
                ],
                "supports": [{"id": "s1", "type": "fixed", "name": "Fixed 1", "faces": [1]}],
                "loads": [
                    {
                        "id": "l1",
                        "type": "pressure",
                        "name": "Pressure 1",
                        "faces": [2],
                        "magnitude": {"value": 10, "unit": "MPa"},
                    }
                ],
                "mesh": {"size": {"value": 5, "unit": "mm"}, "order": 2},
                "solver": "pastix",
            },
            "mesh": {
                "size": {"value": 5, "unit": "mm"},
                "order": 2,
                "nodeCount": 1234,
                "elementCount": 567,
            },
            "software": {
                "clearfea": "0.0.0",
                "deckFormat": 9,
                "calculix": "2.21",
                "gmsh": "4.15.2",
                "occt": "7.9.2",
            },
            "files": ["deck.inp", "solver.log"],
            "machine": {"size": "standard", "ratePerHour": 1000},
            "estimate": None,
            "cost": {"billedSeconds": 3, "amount": 1, "charged": True},
            "status": "succeeded",
            "stage": None,
            "problem": None,
            "result": {
                "nodes": 1234,
                "displacement": {"max": {"value": 1.5e-6, "node": 9, "position": [0.04, 0, 0]}},
                "vonMises": {
                    "min": {"value": 2e6, "node": 1, "position": [0, 0, 0]},
                    "max": {"value": 1.2e7, "node": 2, "position": [0, 0.02, 0]},
                },
                "reactions": [
                    {
                        "supportId": "s1",
                        "force": [-2000.0, 0.0, 0.0],
                        "moment": [0.0, 0.0, 0.0],
                        "point": [0, 0.01, 0.005],
                    }
                ],
                "loads": [{"loadId": "l1", "force": [2000.0, 0, 0], "moment": [0, 0, 0]}],
                "balance": {
                    "appliedForce": [2000.0, 0, 0],
                    "appliedMoment": [0, 0, 0],
                    "reactionForce": [-2000.0, 0, 0],
                    "reactionMoment": [0, 0, 0],
                    "forceScale": 2000,
                    "momentScale": 91.6,
                    "forceRatio": 1.2e-6,
                    "momentRatio": 3e-7,
                },
                "safetyFactor": {"min": {"value": 20.69, "node": 2, "position": [0, 0.02, 0]}},
                "software": {"calculix": "2.21"},
            },
            "createdAt": "2026-10-17T04:00:00.000Z",
            "finishedAt": "2026-10-17T04:00:05.000Z",
        },
        "meshQuality": {
            "shapeMin": 0.31,
            "shapeMean": 0.82,
            "shapeHistogram": [0, 0, 0, 2, 10, 40, 100, 200, 150, 65],
            "jacobianMin": 0.9,
        },
        "figures": [
            {
                "file": "figure-0.png",
                "kind": "support",
                "item": "s1",
                "caption": "Fixed 1 (fixed support)",
                "width": 16,
                "height": 10,
            },
            {
                "file": "figure-1.png",
                "kind": "load",
                "item": "l1",
                "caption": "Pressure 1",
                "width": 16,
                "height": 10,
            },
            {
                "file": "figure-2.png",
                "kind": "view",
                "item": "v1",
                "caption": 'Von Mises stress, view "Overall"',
                "width": 16,
                "height": 10,
            },
        ],
        "statements": [
            {"id": "toolPositioning", "title": "About this report", "text": "General tool."},
            {"id": "softwareVersions", "title": "Software versions", "text": "Versions."},
        ],
        "notes": [],
    }
    params.update(changes)
    return params


def build(params: dict[str, Any], log: str = "") -> Any:
    return build_report(ReportJobParams.model_validate(params), topology(), log)


def texts(document: Any) -> str:
    out = []
    for section in document.sections:
        out.append(section.title)
        for block in section.blocks:
            if isinstance(block, Table):
                out.append(block.caption)
                out.extend(" | ".join(row) for row in [block.columns, *block.rows])
            elif isinstance(block, Paragraph | Figure):
                out.append(block.caption if isinstance(block, Figure) else block.text)
            else:
                out.append(block.text)
    return "\n".join(out)


def test_sections_print_in_order_with_statements_around_them() -> None:
    document = build(static_params())
    titles = [s.title for s in document.sections]
    assert titles == [
        "About this report",
        *[SECTION_TITLES[s] for s in SECTIONS],
        "Software and verification",
    ]


def test_only_chosen_sections_print() -> None:
    document = build(static_params(sections=["materials", "extremes"]))
    assert [s.title for s in document.sections][1:3] == ["Materials", "Minimum and maximum tables"]


def test_values_print_in_the_report_unit_system() -> None:
    text = texts(build(static_params()))
    # Setup values: converted, with the value as entered when typed in another unit.
    assert "Young's modulus | 200000 MPa (entered as 200 GPa)" in text
    assert "Yield strength | 248.2 MPa (entered as 36 ksi)" in text
    assert "Pressure | 10 MPa" in text
    # Results: SI to the system's units.
    assert (
        "Von Mises stress | 2 MPa | node 2 at (0, 0, 0) mm | 12 MPa | node 3 at (0, 20, 0) mm"
        in text
    )
    assert "Fixed 1 | -2000 | 0 | 0 | 0 | 0 | 0" in text
    assert "Fx (N)" in text and "Mx (N·mm)" in text


def test_us_units() -> None:
    text = texts(build(static_params(unitSystem="us-in")))
    assert "Pressure | 1450 psi (entered as 10 MPa)" in text
    assert "Size along x, y, z | (1.575, 0.7874, 0.3937) in" in text
    assert "Fx (lbf)" in text


def test_balance_and_extremes_are_numbers_only() -> None:
    text = texts(build(static_params()))
    assert "Force out of balance, over the loads' size | 0.00012 %" in text
    assert "Safety factor against yield | 20.69 |" in text
    for word in ("PASS", "FAIL", "compliant", "acceptable", "OK"):
        assert re.search(rf"\b{word}\b", text) is None, word


def test_figures_follow_their_items() -> None:
    document = build(static_params())
    figures = document.figures()
    assert [f.caption for f in figures] == [
        "Fixed 1 (fixed support)",
        "Pressure 1",
        'Von Mises stress, view "Overall"',
    ]
    without = build(static_params(figures=[]))
    assert "No figure was drawn for this item." in texts(without)
    assert "it has no result plots" in texts(without)


def test_title_block_cover_and_header() -> None:
    document = build(static_params())
    assert document.header == "Block under pressure · CF-001 · Revision A"
    assert [(line.label, line.value) for line in document.cover] == [
        ("Document number", "CF-001"),
        ("Revision", "A"),
        ("Date", "2026-10-17"),
        ("Prepared by", "A. Engineer"),
    ]
    assert report_date("2026-10-17T23:59:59.000Z") == "2026-10-17"


def test_thermal_run_has_a_heat_balance() -> None:
    params = static_params()
    run = params["run"]
    run["spec"]["analysis"] = "steady_thermal"
    run["spec"]["supports"] = []
    run["spec"]["loads"] = [
        {
            "id": "t1",
            "type": "fixed_temperature",
            "name": "Hot face",
            "faces": [1],
            "temperature": {"value": 100, "unit": "°C"},
        },
        {
            "id": "c1",
            "type": "convection",
            "name": "Air",
            "faces": [2],
            "coefficient": {"value": 10, "unit": "W/(m²·K)"},
            "ambientTemperature": {"value": 20, "unit": "°C"},
        },
    ]
    run["result"] = {
        "nodes": 1234,
        "temperature": {
            "min": {"value": 300.0, "node": 0, "position": [0, 0, 0]},
            "max": {"value": 373.15, "node": 1, "position": [0, 0, 0]},
        },
        "heatFlows": [{"loadId": "t1", "heat": 5.0}, {"loadId": "c1", "heat": -5.0}],
        "heatBalance": {"heatIn": 5.0, "heatOut": 5.0, "scale": 5.0, "ratio": 0.0},
        "software": {"calculix": "2.21"},
    }
    params["figures"] = []
    text = texts(build(params))
    assert "Heat balance check" in text
    assert "Hot face | 5" in text and "Air | -5" in text
    assert "Temperature | 26.85 °C" in text
    assert "Reactions" not in text


def test_modal_run_lists_frequencies() -> None:
    params = static_params()
    params["run"]["spec"]["analysis"] = "modal"
    params["run"]["spec"]["loads"] = []
    params["run"]["spec"]["modes"] = 3
    params["run"]["result"] = {
        "nodes": 1234,
        "frequencies": [64.86, 405.2, 1130.0],
        "software": {"calculix": "2.21"},
    }
    text = texts(build(params))
    assert "Mode | Frequency (Hz)" in text and "2 | 405.2" in text
    assert "A modal analysis applies no loads" in text


def test_notes_print_before_the_closing_statements() -> None:
    note = {"id": "tiedInterfaces", "title": "Stresses near bonded faces", "text": "Ties."}
    document = build(static_params(notes=[note]))
    assert [s.title for s in document.sections][-2:] == [
        "Notes on accuracy",
        "Software and verification",
    ]


def test_solver_warnings_are_listed_once_with_counts() -> None:
    log = " *WARNING in a: one\n x\n *WARNING in a: one\n *WARNING in b: two\n"
    assert solver_warnings(log) == ["*WARNING in a: one (2 times)", "*WARNING in b: two"]
    assert "CalculiX printed no warnings" in texts(build(static_params(), ""))
    assert "*WARNING in b: two" in texts(build(static_params(), log))


def test_a_failed_run_is_refused() -> None:
    params = copy.deepcopy(static_params())
    params["run"]["status"] = "failed"
    with pytest.raises(ValueError, match="succeeded"):
        build(params)


def test_docx_has_the_cover_sections_tables_and_figures(tmp_path: Path) -> None:
    for name in ("figure-0.png", "figure-1.png", "figure-2.png", "logo"):
        (tmp_path / name).write_bytes(png())
    out = tmp_path / "report.docx"
    document = build(static_params())
    placed = write_docx(document, tmp_path, out, "2026-10-17T04:10:00.000Z")
    assert placed == 3
    docx = Document(str(out))
    paragraphs = [p.text for p in docx.paragraphs]
    assert "Block under pressure" in paragraphs
    assert "Example Engineering" in paragraphs
    assert "1  About this report" in paragraphs
    assert "11  Run warnings" in paragraphs
    assert "Figure 2. Pressure 1" in paragraphs
    assert any(p.startswith("Table 1. Project and revision") for p in paragraphs)
    assert len(docx.inline_shapes) == 4  # the logo and three figures
    assert docx.sections[1].header.paragraphs[0].text == document.header
    assert "Page" in docx.sections[1].footer.paragraphs[0].text
    assert docx.core_properties.title == "Block under pressure"


@pytest.mark.parametrize(
    ("labels", "source"),
    [
        (ANALYSIS_KINDS, "simulation.ts"),
        (SUPPORT_KINDS, "simulation.ts"),
        (LOAD_KINDS, "simulation.ts"),
        (SOLVER_LABELS, "estimator.ts"),
    ],
)
def test_labels_match_the_app(labels: dict[str, str], source: str) -> None:
    text = (REPO / "packages/spec/src" / source).read_text(encoding="utf-8")
    for key, label in labels.items():
        assert re.search(rf'\b{key}: "{re.escape(label)}"', text), f"{key}: {label}"


def test_section_titles_match_the_options() -> None:
    text = (REPO / "packages/spec/src/reportOptions.ts").read_text(encoding="utf-8")
    for key, title in SECTION_TITLES.items():
        assert re.search(rf'\b{key}: "{re.escape(title)}', text), f"{key}: {title}"
